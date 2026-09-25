from io import BytesIO
from openpyxl import Workbook, load_workbook
from .models import Category, ImportBatch, ServiceRequest

def safe_cell(value):
    if isinstance(value,str) and value.lstrip().startswith(("=","+","-","@")):
        return "'" + value
    return value

HISTORICAL_HEADERS=["date","reference","name","subject","resolution","handled_by"]

def export_xlsx(queryset):
    wb=Workbook(); ws=wb.active; ws.title="Requests"
    ws.append(["Date","Reference","Requester","Category","Subject","Resolution","Channel","Status","Handled by"])
    for r in queryset.select_related("category"):
        ws.append([safe_cell(v) for v in [r.request_date.isoformat(),r.requester_reference,r.requester_name,r.category.name,r.subject,r.resolution,r.get_channel_display(),r.get_status_display(),r.handled_by]])
    summary=wb.create_sheet("Summary")
    summary.append(["Metric","Value"]); summary.append(["Total",queryset.count()]); summary.append(["Open",queryset.exclude(status__in=["RESOLVED","CANCELLED"]).count()])
    out=BytesIO(); wb.save(out); out.seek(0); return out

def import_history(uploaded_file,user,dry_run=True):
    wb=load_workbook(uploaded_file,data_only=True,read_only=True)
    ws=wb.active
    category,_=Category.objects.get_or_create(name="Imported history")
    created=skipped=0
    for row in ws.iter_rows(min_row=2,values_only=True):
        if not any(row): continue
        values=list(row[:6])+[None]*max(0,6-len(row))
        date,reference,name,subject,resolution,handled_by=values[:6]
        if not date or not name or not subject: skipped+=1; continue
        key=ServiceRequest.build_import_key(values[:6])
        if ServiceRequest.objects.filter(import_key=key).exists(): skipped+=1; continue
        if not dry_run:
            ServiceRequest.objects.create(request_date=date,requester_reference=reference or "",requester_name=name,category=category,subject=subject,resolution=resolution or "",status="RESOLVED" if resolution else "OPEN",handled_by=handled_by or user.get_full_name() or user.username,created_by=user,import_key=key)
        created+=1
    ImportBatch.objects.create(filename=getattr(uploaded_file,"name","history.xlsx"),dry_run=dry_run,created_count=created,skipped_count=skipped,created_by=user)
    return {"created":created,"skipped":skipped,"dry_run":dry_run}
